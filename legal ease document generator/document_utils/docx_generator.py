from io import BytesIO

from docx import Document

from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx.shared import (
    Inches,
    Pt
)

from .common import (
    decode_logo,
    sanitize_text
)


def format_docx(
    text: str,
    doc_type: str,
    logo_base64: str | None = None
) -> bytes:

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal_style = document.styles["Normal"]

    normal_style.font.name = (
        "Times New Roman"
    )

    normal_style.font.size = Pt(11)

    logo = decode_logo(
        logo_base64
    )

    if logo:

        try:

            paragraph = document.add_paragraph()

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )

            paragraph.add_run().add_picture(
                BytesIO(logo),
                width=Inches(1.6)
            )

        except Exception:
            pass

    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title_run = title.add_run(
        doc_type.upper()
    )

    title_run.bold = True

    title_run.font.name = (
        "Times New Roman"
    )

    title_run.font.size = Pt(16)

    clean_text = sanitize_text(text)

    for raw_line in clean_text.split("\n"):

        line = raw_line.strip()

        if not line:

            document.add_paragraph("")

            continue

        upper = line.upper()

        is_heading = (
            len(line) < 90
            and (
                upper == line
                or upper.startswith("SECTION ")
                or upper.startswith("ARTICLE ")
                or upper.startswith("PARTIES")
                or upper.startswith("TERMS")
                or upper.startswith("SIGNATURE")
                or upper.startswith("DRAFT NOTICE")
                or upper.startswith("GENERAL PROVISIONS")
                or upper.startswith("BACKGROUND")
                or upper.startswith("JURISDICTION")
            )
        )

        if is_heading:

            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                line
            )

            run.bold = True

            run.font.name = (
                "Times New Roman"
            )

            run.font.size = Pt(12)

        else:

            paragraph = document.add_paragraph(
                line
            )

            paragraph.paragraph_format.space_after = (
                Pt(5)
            )

            paragraph.paragraph_format.line_spacing = (
                1.15
            )

    footer = section.footer.paragraphs[0]

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer.add_run(
        "LegalEase - AI-generated draft | "
        "Review before legal use"
    )

    footer_run.font.size = Pt(8)

    output = BytesIO()

    document.save(output)

    return output.getvalue()